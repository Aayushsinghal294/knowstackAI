from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer
from sqlalchemy.orm import Session

from app.models.document import Document, DocumentChunk


# Loaded once per backend process.
embedding_model = SentenceTransformer("all-MiniLM-L6-v2")

CHUNK_SIZE = 1000       # characters, not tokens
CHUNK_OVERLAP = 150     # characters
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


def extract_pages(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        import io

        reader = PdfReader(io.BytesIO(data))
        return [
            (page_number, page.extract_text() or "")
            for page_number, page in enumerate(reader.pages, start=1)
        ]

    if extension == ".docx":
        import io

        doc = DocxDocument(io.BytesIO(data))
        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        )
        return [(None, text)]

    if extension in {".txt", ".md"}:
        text = data.decode("utf-8-sig")
        return [(None, text)]

    raise ValueError("Unsupported file type. Use PDF, DOCX, TXT or MD.")


def split_text(text: str) -> list[str]:
    text = " ".join(text.split())

    if not text:
        return []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end == len(text):
            break

        start = end - CHUNK_OVERLAP

    return chunks



def ingest_document(
    db: Session,
    filename: str,
    data: bytes,
    *,
    user_id: int,
) -> dict:
    if len(data) > MAX_FILE_SIZE:
        raise ValueError("File exceeds the 10 MB limit.")

    pages = extract_pages(filename, data)

    prepared_chunks = []

    for page_number, page_text in pages:
        for chunk in split_text(page_text):
            prepared_chunks.append((page_number, chunk))

    if not prepared_chunks:
        raise ValueError("No readable text found in the document.")

    texts = [chunk for _, chunk in prepared_chunks]

    vectors = embedding_model.encode(
        texts,
        normalize_embeddings=True,
        show_progress_bar=False,
    ).tolist()

    try:
        document = Document(
            user_id=user_id,
            filename=Path(filename).name,
            file_type=Path(filename).suffix.lower().lstrip("."),
            status="processed",
        )

        db.add(document)
        db.flush()

        for index, ((page_number, content), vector) in enumerate(
            zip(prepared_chunks, vectors)
        ):
            db.add(
                DocumentChunk(
                    document_id=document.id,
                    chunk_index=index,
                    page_number=page_number,
                    content=content,
                    embedding=vector,
                )
            )

        db.commit()
        db.refresh(document)

        return {
            "id": document.id,
            "filename": document.filename,
            "chunks_created": len(prepared_chunks),
            "status": document.status,
        }

    except Exception:
        db.rollback()
        raise
